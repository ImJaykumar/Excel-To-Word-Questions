import io
import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor
import pandas as pd
import streamlit as st

# ==========================================
# PAGE CONFIGURATION
# ==========================================
st.set_page_config(
    page_title="Excel to 2-Column Word Converter",
    page_icon="📄",
    layout="wide",
    initial_sidebar_state="expanded",
)

# Custom CSS for Streamlit UI Styling
st.markdown(
    """
    
    """,
    unsafe_allow_html=True,
)


# ==========================================
# WORD CONVERTER HELPER FUNCTIONS
# ==========================================
def set_cell_shading(cell, color_hex):
    shading = parse_xml(f'')
    cell._tc.get_or_add_tcPr().append(shading)


def enable_two_columns(section, col_space_pt=14):
    """Enables 2-column layout on a Word document section."""
    sectPr = section._sectPr
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sectPr.append(cols)
    cols.set(qn("w:num"), "2")
    cols.set(
        qn("w:space"), str(int(col_space_pt * 20))
    )  # Space between columns in dxa


def convert_df_to_two_column_word(df, main_title="CURRENT AFFAIRS & MCQS", sub_title="Practice Questions & Explanations"):
    """Converts DataFrame to a 2-Column Word Document and returns BytesIO buffer."""
    doc = docx.Document()

    # 1. Compact Page Setup (0.5 inch margins)
    section = doc.sections[0]
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    # 2. Enable 2 Columns
    enable_two_columns(section, col_space_pt=14)

    # Color Palette Definitions
    NAVY = RGBColor(15, 30, 54)  # #0F1E36 - Header
    TEAL = RGBColor(13, 148, 136)  # #0D9488 - Accent
    DARK_TEXT = RGBColor(30, 41, 59)  # #1E293B - Regular Option Text
    GREEN_TEXT = RGBColor(22, 101, 52)  # #166534 - Correct Option Text
    AMBER_DARK = RGBColor(180, 83, 9)  # #B45309 - Answer/Explanation Text

    HEX_NAVY = "0F1E36"
    HEX_LIGHT_BG = "F8FAFC"
    HEX_AMBER_BG = "FFFBEB"
    HEX_BORDER_LIGHT = "E2E8F0"
    HEX_AMBER_BORDER = "FCD34D"

    # Header Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run(main_title)
    run_title.font.name = "Segoe UI"
    run_title.font.size = Pt(13)
    run_title.font.bold = True
    run_title.font.color.rgb = NAVY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(8)
    run_sub = p_sub.add_run(sub_title)
    run_sub.font.name = "Segoe UI"
    run_sub.font.size = Pt(8.5)
    run_sub.font.bold = True
    run_sub.font.color.rgb = TEAL

    COL_WIDTH = Inches(3.55)

    for index, row in df.iterrows():
        q_num = row.get("Q. No.", index + 1)
        question_text = str(row.get("Question", "")).strip()

        opt_a = str(row.get("Option A", "")).strip() if pd.notna(row.get("Option A")) else ""
        opt_b = str(row.get("Option B", "")).strip() if pd.notna(row.get("Option B")) else ""
        opt_c = str(row.get("Option C", "")).strip() if pd.notna(row.get("Option C")) else ""
        opt_d = str(row.get("Option D", "")).strip() if pd.notna(row.get("Option D")) else ""

        correct_ans = str(row.get("Correct Answer", "")).strip() if pd.notna(row.get("Correct Answer")) else ""
        explanation = str(row.get("Explanation", "")).strip() if pd.notna(row.get("Explanation")) else ""

        # Question Box
        table_q = doc.add_table(rows=1, cols=1)
        table_q.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_q.autofit = False
        cell_q = table_q.cell(0, 0)
        cell_q.width = COL_WIDTH

        set_cell_shading(cell_q, HEX_LIGHT_BG)

        tcPr = cell_q._tc.get_or_add_tcPr()
        borders = parse_xml(
            f''
            f''
            f''
            f''
            f''
            f""
        )
        tcPr.append(borders)

        tcMar = OxmlElement("w:tcMar")
        for m, val in [("top", 80), ("bottom", 80), ("left", 100), ("right", 100)]:
            node = OxmlElement(f"w:{m}")
            node.set(qn("w:w"), str(val))
            node.set(qn("w:type"), "dxa")
            tcMar.append(node)
        tcPr.append(tcMar)

        p_q = cell_q.paragraphs[0]
        p_q.paragraph_format.space_before = Pt(1)
        p_q.paragraph_format.space_after = Pt(1)
        p_q.paragraph_format.line_spacing = 1.05

        run_qbadge = p_q.add_run(f"Q{q_num}. ")
        run_qbadge.font.name = "Segoe UI"
        run_qbadge.font.size = Pt(9.5)
        run_qbadge.font.bold = True
        run_qbadge.font.color.rgb = TEAL

        run_qbody = p_q.add_run(question_text)
        run_qbody.font.name = "Segoe UI"
        run_qbody.font.size = Pt(9.5)
        run_qbody.font.bold = True
        run_qbody.font.color.rgb = NAVY

        # 2x2 Options Grid
        options_list = [("A", opt_a), ("B", opt_b), ("C", opt_c), ("D", opt_d)]

        table_opt = doc.add_table(rows=2, cols=2)
        table_opt.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_opt.autofit = False

        for row_idx in range(2):
            for col_idx in range(2):
                opt_key, opt_val = options_list[row_idx * 2 + col_idx]
                cell_opt = table_opt.cell(row_idx, col_idx)
                cell_opt.width = Inches(1.75)

                is_correct = (correct_ans.upper() == opt_key.upper()) or (
                    f"OPTION {opt_key}".upper() in correct_ans.upper()
                )

                set_cell_shading(cell_opt, "FFFFFF")

                tcPr_opt = cell_opt._tc.get_or_add_tcPr()
                opt_borders = parse_xml(
                    f''
                    f''
                    f''
                    f''
                    f''
                    f""
                )
                tcPr_opt.append(opt_borders)

                tcMar_opt = OxmlElement("w:tcMar")
                for m, val in [("top", 60), ("bottom", 60), ("left", 80), ("right", 80)]:
                    node = OxmlElement(f"w:{m}")
                    node.set(qn("w:w"), str(val))
                    node.set(qn("w:type"), "dxa")
                    tcMar_opt.append(node)
                tcPr_opt.append(tcMar_opt)

                p_opt = cell_opt.paragraphs[0]
                p_opt.paragraph_format.space_before = Pt(1)
                p_opt.paragraph_format.space_after = Pt(1)
                p_opt.paragraph_format.line_spacing = 1.05

                run_badge = p_opt.add_run(f"({opt_key}) ")
                run_badge.font.name = "Calibri"
                run_badge.font.size = Pt(9)
                run_badge.font.bold = True
                run_badge.font.color.rgb = GREEN_TEXT if is_correct else TEAL

                run_text = p_opt.add_run(opt_val)
                run_text.font.name = "Calibri"
                run_text.font.size = Pt(9)
                run_text.font.color.rgb = GREEN_TEXT if is_correct else DARK_TEXT
                if is_correct:
                    run_text.font.bold = True

        # Explanation Box
        if correct_ans or explanation:
            table_exp = doc.add_table(rows=1, cols=1)
            table_exp.alignment = WD_TABLE_ALIGNMENT.CENTER
            table_exp.autofit = False
            cell_exp = table_exp.cell(0, 0)
            cell_exp.width = COL_WIDTH

            set_cell_shading(cell_exp, HEX_AMBER_BG)

            tcPr_exp = cell_exp._tc.get_or_add_tcPr()
            exp_borders = parse_xml(
                f''
                f''
                f''
                f''
                f''
                f""
            )
            tcPr_exp.append(exp_borders)

            tcMar_exp = OxmlElement("w:tcMar")
            for m, val in [("top", 60), ("bottom", 60), ("left", 80), ("right", 80)]:
                node = OxmlElement(f"w:{m}")
                node.set(qn("w:w"), str(val))
                node.set(qn("w:type"), "dxa")
                tcMar_exp.append(node)
            tcPr_exp.append(tcMar_exp)

            p_exp = cell_exp.paragraphs[0]
            p_exp.paragraph_format.space_before = Pt(1)
            p_exp.paragraph_format.space_after = Pt(1)
            p_exp.paragraph_format.line_spacing = 1.05

            if correct_ans:
                run_ans_lbl = p_exp.add_run("✔ Ans: ")
                run_ans_lbl.font.name = "Segoe UI"
                run_ans_lbl.font.size = Pt(8.5)
                run_ans_lbl.font.bold = True
                run_ans_lbl.font.color.rgb = AMBER_DARK

                run_ans_val = p_exp.add_run(f"Option {correct_ans}\n")
                run_ans_val.font.name = "Segoe UI"
                run_ans_val.font.size = Pt(8.5)
                run_ans_val.font.bold = True
                run_ans_val.font.color.rgb = AMBER_DARK

            if explanation:
                run_exp_lbl = p_exp.add_run("💡 ")
                run_exp_lbl.font.name = "Segoe UI"
                run_exp_lbl.font.size = Pt(8.5)
                run_exp_lbl.font.bold = True
                run_exp_lbl.font.color.rgb = AMBER_DARK

                run_exp_val = p_exp.add_run(explanation)
                run_exp_val.font.name = "Calibri"
                run_exp_val.font.size = Pt(8.5)
                run_exp_val.font.color.rgb = DARK_TEXT

        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_before = Pt(0)
        p_space.paragraph_format.space_after = Pt(6)

    # Save to memory stream
    output_stream = io.BytesIO()
    doc.save(output_stream)
    output_stream.seek(0)
    return output_stream


# ==========================================
# STREAMLIT UI LAYOUT
# ==========================================
st.title("📄 Excel / CSV to 2-Column Word Converter")
st.caption("MCQ प्रश्नों की Excel/CSV फ़ाइल अपलोड करें और उसे प्रीमियम 2-कॉलम वर्ड डॉक्यूमेंट (.docx) में कनवर्ट करें।")

st.sidebar.title("⚙️ Document Settings")
main_title = st.sidebar.text_input("मुख्य शीर्षक (Main Title)", "CURRENT AFFAIRS & MCQS")
sub_title = st.sidebar.text_input("उप-शीर्षक (Sub Title)", "Practice Questions & Explanations")

uploaded_file = st.file_uploader("Excel (.xlsx) या CSV (.csv) फ़ाइल चुनें", type=["xlsx", "xls", "csv"])

if uploaded_file is not None:
    try:
        if uploaded_file.name.endswith(".csv"):
            df = pd.read_csv(uploaded_file)
        else:
            df = pd.read_excel(uploaded_file)

        st.success(f"✅ फ़ाइल सफलतापूर्वक लोड हुई! कुल {len(df)} प्रश्न पाए गए।")

        # Preview Data
        with st.expander("👀 डेटा पूर्वावलोकन (Preview First 5 Rows)"):
            st.dataframe(df.head(5), use_container_width=True)

        if st.button("🚀 2-Column Word Document कनवर्ट करें"):
            with st.spinner("2-Column Word फ़ाइल तैयार की जा रही है..."):
                docx_buffer = convert_df_to_two_column_word(df, main_title, sub_title)
                
                output_filename = f"{uploaded_file.name.rsplit('.', 1)[0]}_2Col_Formatted.docx"
                
                st.download_button(
                    label="📥 Download 2-Column Word Document (.docx)",
                    data=docx_buffer,
                    file_name=output_filename,
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
    except Exception as e:
        st.error(f"❌ फ़ाइल प्रोसेस करने में त्रुटि: {e}")
